"""Build an editable Word report with real footnotes from the assessment Markdown."""
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parent
TEXT = (ROOT / "CCWC_Research_Assessment.md").read_text()
SOURCES = dict(re.findall(r"^\[\^(\d+)\]: (.+)$", TEXT, re.M))
BODY = TEXT.split("## Sources\n", 1)[0]
ORDER = list(dict.fromkeys(re.findall(r"\[\^(\d+)\]", BODY)))
NUMBERS = {key: i + 1 for i, key in enumerate(ORDER)}
TOKEN = re.compile(r"(\[\^\d+\]|\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def xml_run(text, bold=False, superscript=False, size=None):
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    if bold:
        props.append(OxmlElement("w:b"))
    if superscript:
        v = OxmlElement("w:vertAlign")
        v.set(qn("w:val"), "superscript")
        props.append(v)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size * 2))
        props.append(sz)
    run.append(props)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    return run


def hyperlink(parent, part, label, url=None, anchor=None, superscript=False):
    link = OxmlElement("w:hyperlink")
    if url:
        link.set(qn("r:id"), part.relate_to(url, RT.HYPERLINK, is_external=True))
    else:
        link.set(qn("w:anchor"), anchor)
    link.append(xml_run(label, superscript=superscript))
    parent.append(link)


def plain_inline(parent, part, text):
    for token in TOKEN.split(text):
        if not token:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if match:
            hyperlink(parent, part, match[1], url=match[2])
        elif token.startswith("**") and token.endswith("**"):
            parent.append(xml_run(token[2:-2], bold=True))
        elif token.startswith("`"):
            parent.append(xml_run(token[1:-1]))
        else:
            parent.append(xml_run(token))


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    for name in ("Footnote Text", "Source Text"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.base_style = normal
        style.font.size = Pt(9)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.0

    notes = OxmlElement("w:footnotes")
    for note_id, kind in ((-1, "separator"), (0, "continuationSeparator")):
        note = OxmlElement("w:footnote")
        note.set(qn("w:type"), kind)
        note.set(qn("w:id"), str(note_id))
        para = OxmlElement("w:p")
        run = OxmlElement("w:r")
        run.append(OxmlElement("w:" + kind))
        para.append(run)
        note.append(para)
        notes.append(note)
    footpart = Part(PackURI("/word/footnotes.xml"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                    b"", doc.part.package)
    doc.part.relate_to(footpart, RT.FOOTNOTES)
    used = set()

    def inline(paragraph, content):
        for token in TOKEN.split(content):
            match = re.fullmatch(r"\[\^(\d+)\]", token)
            if not match:
                plain_inline(paragraph._p, doc.part, token)
                continue
            key = match[1]
            number = NUMBERS[key]
            if key in used:
                hyperlink(paragraph._p, doc.part, str(number), anchor=f"source_{number}", superscript=True)
                continue
            used.add(key)
            run = paragraph.add_run()
            run.font.superscript = True
            ref = OxmlElement("w:footnoteReference")
            ref.set(qn("w:id"), str(number))
            run._r.append(ref)
            note = OxmlElement("w:footnote")
            note.set(qn("w:id"), str(number))
            para = OxmlElement("w:p")
            props = OxmlElement("w:pPr")
            sty = OxmlElement("w:pStyle")
            sty.set(qn("w:val"), doc.styles["Footnote Text"].style_id)
            props.append(sty)
            para.append(props)
            marker = OxmlElement("w:r")
            marker.append(OxmlElement("w:footnoteRef"))
            para.append(marker)
            para.append(xml_run(" "))
            plain_inline(para, footpart, SOURCES[key])
            note.append(para)
            notes.append(note)

    lines = BODY.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        if not line.strip():
            continue
        if line.startswith("| "):
            rows = [line]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            parsed = [[c.strip() for c in row.strip("|").split("|")] for row in rows]
            parsed = [r for r in parsed if not all(re.fullmatch(r":?-+:?", c) for c in r)]
            table = doc.add_table(rows=0, cols=len(parsed[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = {2: [2.0, 5.1], 3: [1.35, 3.45, 2.3]}.get(len(parsed[0]))
            for ri, row in enumerate(parsed):
                cells = table.add_row().cells
                if widths:
                    for cell, width in zip(cells, widths):
                        cell.width = Inches(width)
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                for cell, content in zip(cells, row):
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.space_before = Pt(3)
                    para.paragraph_format.line_spacing = 1.0
                    inline(para, content)
                    for run in para.runs:
                        run.font.size = Pt(9)
                        if ri == 0:
                            run.bold = True
                    if ri == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "EEEEEE")
                        cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif line.startswith("# "):
            inline(doc.add_paragraph(style="Title"), line[2:])
        elif line.startswith("## "):
            inline(doc.add_paragraph(style="Heading 1"), line[3:])
        elif line.startswith("### "):
            inline(doc.add_paragraph(style="Heading 2"), line[4:])
        elif re.match(r"^\d+\. ", line):
            inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        else:
            inline(doc.add_paragraph(), line)

    doc.add_heading("Sources", level=1)
    for key in ORDER:
        number = NUMBERS[key]
        para = doc.add_paragraph(style="Source Text")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(number))
        start.set(qn("w:name"), f"source_{number}")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(number))
        para._p.append(start)
        plain_inline(para._p, doc.part, f"{number}. {SOURCES[key]}")
        para._p.append(end)
    footpart._blob = etree.tostring(notes, xml_declaration=True, encoding="UTF-8", standalone=True)
    doc.core_properties.title = "Solar Forecasting Research Assessment for IEEE CCWC"
    doc.core_properties.subject = "Research audit, contribution selection, and experiment plan"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "solar forecasting, CCWC, research assessment"
    path = ROOT / "CCWC_Research_Assessment.docx"
    doc.save(path)
    print(f"Wrote {path.name}; {len(used)} linked footnotes; {len(doc.tables)} tables")


if __name__ == "__main__":
    build()
