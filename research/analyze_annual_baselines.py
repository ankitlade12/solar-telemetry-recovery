"""Development tables, event-level paired uncertainty and standalone figures."""
import argparse
from importlib.metadata import version
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from solar_recovery.pilot import digest

PHASES = ["outage", "partial_restoration", "recovery_0_6", "recovery_6_24"]


def table(frame, digits=4):
    lines = ["| "+" | ".join(map(str, frame.columns))+" |", "| "+" | ".join("---" for _ in frame)+" |"]
    for values in frame.itertuples(index=False, name=None):
        lines.append("| "+" | ".join(f"{v:.{digits}f}" if isinstance(v, float) else str(v) for v in values)+" |")
    return "\n".join(lines)


def paired_events(left, right, contrast):
    keys = ["origin", "horizon", "event_id", "phase"]
    if "scenario" in left and "scenario" in right:
        keys.append("scenario")
    merged = left[keys+["nwis"]].merge(right[keys+["nwis"]], on=keys, suffixes=("_left", "_right"), validate="one_to_one")
    merged["difference"] = merged.nwis_left-merged.nwis_right
    by_event_scenario = merged.groupby(["phase", "event_id"]+(["scenario"] if "scenario" in keys else [])).difference.mean()
    by_event = by_event_scenario.groupby(level=["phase", "event_id"]).mean().reset_index()
    by_event["contrast"] = contrast
    return by_event


def analyze(root):
    root = Path(root)
    manifest = json.loads((root/"manifest.json").read_text())
    assert manifest["status"].startswith("complete")
    selected = pd.read_csv(root/"selected_profiles.csv")
    metrics = pd.read_csv(root/"metrics.csv")
    all_summary, phase_summary, event_differences = [], [], []
    for site in selected.site.unique():
        choices = selected.loc[selected.site.eq(site)].method.tolist()+["persistence", "geometric_persistence"]
        overall = pd.read_csv(root/site/"overall_metrics.csv")
        development = overall.loc[overall.period.eq("development_evaluation") & overall.method.isin(choices)]
        summary = development.groupby("method")[["nwis", "nmae", "coverage50", "coverage80", "coverage90", "nwidth90"]].mean().reset_index()
        summary.insert(0, "site", site)
        all_summary.append(summary)
        phases = metrics.loc[metrics.site.eq(site) & metrics.period.eq("development_evaluation") & metrics.method.isin(choices)]
        phase_summary.append(phases.groupby(["site", "phase", "method"])[["nwis", "coverage90", "nwidth90"]].mean().reset_index())
        records = []
        columns = ["origin", "horizon", "method", "event_id", "phase", "nwis", "valid_target", "daylight", "within_period", "period", "scenario"]
        for path in sorted((root/site).glob("forecasts_*.parquet")):
            scored = pd.read_parquet(path, columns=columns)
            scored = scored.loc[scored.valid_target & scored.daylight & scored.within_period & scored.period.eq("development_evaluation") & scored.phase.isin(PHASES)]
            records.append(scored)
        scored = pd.concat(records, ignore_index=True)
        for profile in ("compact", "flexible"):
            left = scored.loc[scored.method.eq(f"multiple_imputation_{profile}")]
            right = scored.loc[scored.method.eq(f"single_imputation_{profile}")]
            event_differences.append(paired_events(left, right, f"MI_minus_SI_{profile}").assign(site=site))
        for method in choices:
            baseline = scored.loc[scored.scenario.eq("joint_immediate") & scored.method.eq(method)].drop(columns="scenario")
            for scenario in ("joint_gradual", "joint_none"):
                candidate = scored.loc[scored.scenario.eq(scenario) & scored.method.eq(method)].drop(columns="scenario")
                event_differences.append(paired_events(candidate, baseline, f"{scenario}_minus_immediate__{method}").assign(site=site))
    summary = pd.concat(all_summary, ignore_index=True)
    phases = pd.concat(phase_summary, ignore_index=True)
    differences = pd.concat(event_differences, ignore_index=True)
    bootstrap = []
    rng = np.random.default_rng(20260911)
    for (site, phase, contrast), group in differences.groupby(["site", "phase", "contrast"]):
        values = group.difference.to_numpy()
        draws = rng.choice(values, size=(10000, len(values)), replace=True).mean(axis=1)
        lower, upper = np.quantile(draws, [.025, .975])
        bootstrap.append({"site": site, "phase": phase, "contrast": contrast, "events": len(values),
            "mean_nwis_difference": float(values.mean()), "ci_lower": float(lower), "ci_upper": float(upper)})
    bootstrap = pd.DataFrame(bootstrap)
    for name, frame in (("site_method_summary", summary), ("phase_summary", phases),
                        ("paired_event_differences", differences), ("event_bootstrap", bootstrap)):
        frame.to_csv(root/f"{name}.csv", index=False)
    figures = root/"figures"
    figures.mkdir(exist_ok=True)
    plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False})
    fig, axes = plt.subplots(2, 2, figsize=(10, 7), constrained_layout=True)
    for row, site in enumerate(selected.site.unique()):
        part = summary.loc[summary.site.eq(site)].sort_values("nwis")
        names = part.method.str.replace("_", " ").str.replace("multiple imputation", "MI").str.replace("single imputation", "SI")
        axes[row, 0].barh(names, part.nwis, color="#315b82")
        axes[row, 0].invert_yaxis()
        axes[row, 0].set(title=f"{site.title()}: score (lower is better)", xlabel="Normalized WIS")
        axes[row, 1].barh(names, part.coverage90*100, color="#327e71")
        axes[row, 1].invert_yaxis()
        axes[row, 1].axvline(90, color="#a43b39", linestyle="--", linewidth=1)
        axes[row, 1].set(title=f"{site.title()}: nominal 90% intervals", xlabel="Empirical coverage (%)", xlim=(0, 100))
    fig.suptitle("July–December 2016 development: equal scenario/horizon averages")
    for extension in ("png", "pdf"):
        fig.savefig(figures/f"annual_baselines.{extension}", dpi=200)
    plt.close(fig)
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.4), constrained_layout=True)
    short_phases = ["recovery_0_6", "recovery_6_24"]
    for ax, site in zip(axes, selected.site.unique()):
        part = bootstrap.loc[bootstrap.site.eq(site) & bootstrap.contrast.eq("MI_minus_SI_compact") & bootstrap.phase.isin(short_phases)].set_index("phase").reindex(short_phases)
        value = part.mean_nwis_difference.to_numpy()
        ax.errorbar([0, 1], value, yerr=np.maximum(0., np.vstack([value-part.ci_lower, part.ci_upper-value])), fmt="o", capsize=4, color="#315b82")
        ax.axhline(0., color="gray", linestyle="--")
        ax.set(xticks=[0, 1], xticklabels=["Recovery 0–6 h", "Recovery 6–24 h"],
            title=site.title(), ylabel="MI minus SI: paired event mean NWIS")
    fig.suptitle("Matched compact models; exploratory 95% event bootstrap intervals")
    for extension in ("png", "pdf"):
        fig.savefig(figures/f"imputation_event_comparison.{extension}", dpi=200)
    plt.close(fig)
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Readme"
    for row in [
        ["Purpose", "Source tables for annual development figures; no final-evaluation claims"],
        ["Run", str(root)], ["Manifest SHA256", digest(root/"manifest.json")],
        ["Main averages", "Equal weight to available scenario/horizon cells within site; daylight observed targets only"],
        ["Event intervals", "10000 paired event resamples, seed 20260911, within site; no multiplicity correction"],
        ["Event weighting", "Average forecasts within event/scenario then average scenarios within event; equal event weight"],
        ["Inference scope", "Conditional on these sites and one calendar schedule; seven-day event spacing does not establish independence"],
        ["Training scope", "Daylight-only targets; nighttime forecasts set to zero and outside scoring scope"],
        ["MI source", "https://arxiv.org/abs/2603.15564 (Setup-2 adaptation, not reproduction)"],
        ["Data source", "https://doi.org/10.25984/1846021"],
    ]:
        readme.append(row)
    for name, frame in (("SelectedProfiles", selected), ("SiteMethodSummary", summary), ("PhaseSummary", phases),
                        ("EventBootstrap", bootstrap), ("EventDifferences", differences)):
        sheet = workbook.create_sheet(name)
        for row in dataframe_to_rows(frame, index=False, header=True):
            sheet.append([None if isinstance(v, float) and not np.isfinite(v) else v for v in row])
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315B82")
        for cells in sheet.columns:
            sheet.column_dimensions[cells[0].column_letter].width = min(65, max(14, max(len(str(c.value or "")) for c in cells)+2))
    workbook.save(root/"Annual_Development_Evidence.xlsx")
    reopened = load_workbook(root/"Annual_Development_Evidence.xlsx", read_only=True)
    assert reopened["SiteMethodSummary"].max_row == len(summary)+1
    reopened.close()
    notable = bootstrap.loc[bootstrap.contrast.eq("MI_minus_SI_compact") & bootstrap.phase.isin(short_phases)]
    note = """# Annual AC development findings

The annual comparison is development evidence, not a final test or an established recovery-method contribution. It uses two geographic sites, 2015 training, February–June 2016 profile selection, and July–December 2016 diagnostics. A prior all-hour fit produced degenerate NIST quantile models and is excluded from claims of algorithmic superiority. This corrected run trains all baseline families on daylight targets and makes stochastic imputations invariant to future rows and batch order.

## Main comparison

Entries below average equally across eight scenarios and four horizons. NWIS and width are normalized by documented DC nameplate capacity. Coverage is a fraction; the nominal level is 0.90. Lower WIS alone does not establish calibrated uncertainty.

"""+table(summary[["site", "method", "nwis", "nmae", "coverage90", "nwidth90"]])+"""

![Annual development comparison](figures/annual_baselines.png)

## Matched imputation comparison

These differences compare ten stochastic test-input completions with deterministic mean completion under identical compact fitted point models. Negative differences favor MI. Each event averages its eligible forecast differences within scenario, then averages scenarios. The percentile bootstrap resamples event IDs within a site; intervals are exploratory, unadjusted for multiple comparisons and conditional on this schedule. They cannot establish generalization across geographic populations or unobserved outage processes.

"""+table(notable[["site", "phase", "events", "mean_nwis_difference", "ci_lower", "ci_upper"]])+"""

![Event comparison](figures/imputation_event_comparison.png)

Backfill contrasts for each selected model, full phase tables and per-event differences are in the CSV files and `Annual_Development_Evidence.xlsx`. The compact and flexible SI/MI pairs are retained even when another profile is selected. All methods share observed evaluation targets; naturally absent targets remain unscored.

## What this establishes and what remains

The experiment establishes a reproducible two-site comparison with separate live resumption, immediate backlog, gradual backlog and permanent-loss schedules. It does not yet establish a recovery-aware calibration improvement. The published imputation comparator is adapted to joint-channel missingness, observed-only training labels, extra causal features and gradient boosting; it is not a faithful reproduction of the original model/tuning configuration. Training-residual normal intervals have no coverage guarantee.

Complete matched physical-context, mask, recency and delayed-adaptive calibration; repeated outage seeds and durations; timestamp/quality sensitivity; and the frozen 2017 evaluation before finalizing a regular-paper claim. The original PRD and all prior run artifacts remain preserved. See [the full development protocol](../../research/ANNUAL_DEVELOPMENT_PROTOCOL.md) for data contracts, timing, method choices and primary sources.
"""
    (root/"FINDINGS.md").write_text(note)
    (root/"analysis_manifest.json").write_text(json.dumps({"analyzer_sha256": digest(__file__), "run_manifest_sha256": digest(root/"manifest.json"),
        "versions": {name: version(name) for name in ("numpy", "pandas", "matplotlib", "openpyxl")},
        "bootstrap_replicates": 10000, "bootstrap_seed": 20260911, "status": "exploratory development analysis"}, indent=2)+"\n")
    print(table(summary[["site", "method", "nwis", "coverage90"]]), flush=True)
    print(f"Wrote findings, CSVs, figures and verified spreadsheet under {root}", flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", required=True)
    args = parser.parse_args()
    analyze(args.run)
