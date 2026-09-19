"""Build the supporting spreadsheet for the development research report."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import pandas as pd


def build():
    workbook = Workbook()
    workbook.remove(workbook.active)
    note = workbook.create_sheet("ReadMe")
    for row in [
        ["Development evidence", "Interpretation"],
        ["Status", "Exploratory development results; no final-test, novelty or cross-site performance claim."],
        ["RunMetrics", "Forecast-weighted phase/split aggregates from the frozen 2015 DC-target experiment."],
        ["EventComparisons", "Figure data: equal-event paired differences and exploratory 95% percentile intervals."],
        ["EventDifferences", "Individual paired event means used for 5,000 bootstrap resamples; only eight events per slice."],
        ["DataQuality", "New 2016 AC-target qualification counts. These are not the target/data used by RunMetrics."],
        ["Reproduction", "python3 -m research.build_development_workbook"],
        ["Reference", "research/DEVELOPMENT_V2_FINDINGS.md and adjacent CSV/JSON source artifacts."]]:
        note.append(row)
    run = Path("runs/nist_2015_calibration_v2")
    for sheet, filename in [("RunMetrics", "metrics.csv"), ("EventComparisons", "event_comparisons.csv"),
                            ("EventDifferences", "event_differences.csv")]:
        frame = pd.read_csv(run/filename)
        tab = workbook.create_sheet(sheet)
        tab.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            tab.append([None if pd.isna(value) else value for value in row])
    quality = workbook.create_sheet("DataQuality")
    quality.append(["contract", "system_id", "year", "hours", "valid_pv_hours", "daylight_hours",
                    "valid_daylight_joint_hours", "joint_daylight_fraction", "source_audit"])
    for name in ["nist_ac_2016", "colorado_ac_2016"]:
        path = Path("data/processed")/name/"audit.json"
        audit = json.loads(path.read_text())
        quality.append([name, audit["system_id"], audit["year"], audit["hours"], audit["valid_pv_hours"],
                        audit["daylight_hours"], audit["valid_daylight_joint_hours"],
                        audit["valid_daylight_joint_hours"]/audit["daylight_hours"], str(path)])
    sources = workbook.create_sheet("Sources")
    sources.append(["source", "title", "year", "URL or artifact", "used for"])
    for row in [
        ["IEEE CCWC", "Submissions", 2027, "https://ieee-ccwc.org/submissions/", "Deadline and anonymity"],
        ["IEEE CCWC", "Call for papers", 2027, "https://ieee-ccwc.org/call-for-papers/", "Paper category and page limit"],
        ["NIST", "Photovoltaic data dictionary", None, "https://www.nist.gov/document/datadictionarysupplementalcontentpdf", "Instrument definitions"],
        ["Boyd", "Performance Data from the NIST Photovoltaic Arrays and Weather Station", 2017, "https://doi.org/10.6028/jres.122.040", "Time and aggregation conventions"],
        ["Romano et al.", "Conformalized Quantile Regression", 2019, "https://arxiv.org/abs/1905.03222", "Calibration reference"],
        ["Gibbs and Candes", "Adaptive Conformal Inference Under Distribution Shift", 2021, "https://arxiv.org/abs/2106.00170", "Adaptive update reference"],
        ["Local experiment", "Calibration v2 forecast/metric artifacts", 2026, str(run), "All reported forecasting numbers"],
        ["Local audit", "PVDAQ development qualification", 2026, "research/development_source_audit.json", "Channel and timestamp availability"]]:
        sources.append(row)
    for tab in workbook:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="333333")
        for column in tab.columns:
            length = min(65, max(14, max(len(str(cell.value or "")) for cell in column)+2))
            tab.column_dimensions[column[0].column_letter].width = length
        for row in tab.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=tab.title in {"ReadMe", "Sources"})
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
    output = Path("research/Development_Evidence.xlsx")
    workbook.save(output)
    check = load_workbook(output, read_only=True, data_only=True)
    assert check["EventComparisons"].max_row == 7
    assert check["DataQuality"].max_row == 3
    print(f"Saved and reopened {output}: {check.sheetnames}")


if __name__ == "__main__":
    build()
