"""Separate producing-system analysis: unchanged estimands, single-site figures, authenticated extension inputs."""
import argparse
import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from threadpoolctl import threadpool_limits

from solar_recovery.pilot import digest

SCORES = ["nwis", "nmae", "coverage50", "coverage80", "coverage90", "nwidth90"]
PHASES = {
    "outage": ["outage"], "partial_restoration": ["partial_restoration"],
    "recovery_0_6": ["recovery_0_6"], "recovery_6_24": ["recovery_6_24"],
    "recovery": ["recovery_0_6", "recovery_6_24"],
    "failure_recovery": ["outage", "partial_restoration", "recovery_0_6", "recovery_6_24"],
    "all": None,
}
INTERRUPTION_CASES = {"clean", "pv_only", "irradiance_only", "joint_immediate", "joint_gradual"}


def interruption_comparison(by_site):
    """Raw-forecast descriptive curves on the same scheduled-interruption rows."""
    results = []
    for site, cases in by_site.items():
        if set(cases) != INTERRUPTION_CASES:
            raise ValueError("Interruption curves require clean, isolated and joint loss cases")
        indexed = {name: frame.set_index(["origin", "horizon"]).sort_index() for name, frame in cases.items()}
        reference = indexed["joint_gradual"]
        reference = reference.loc[reference.phase.isin(["outage", "partial_restoration"])]
        for case, frame in indexed.items():
            if case != "clean":
                interrupted = frame.loc[frame.phase.isin(["outage", "partial_restoration"])]
                if not interrupted.index.equals(reference.index):
                    raise ValueError("Interruption cases do not share the same scheduled scoring window")
            if not reference.index.isin(frame.index).all():
                raise ValueError("Missing common interruption scoring rows")
            matched = frame.loc[reference.index].copy()
            if not matched.zero_with_bright_poa.equals(reference.zero_with_bright_poa):
                raise ValueError("Interruption comparison has inconsistent target-ambiguity flags")
            matched["reference_event_id"] = reference.event_id.to_numpy()
            for population in ("recorded_ac", "exclude_bright_zero"):
                selected = matched if population == "recorded_ac" else matched.loc[~matched.zero_with_bright_poa]
                values = selected.groupby("horizon").agg(n=("nwis", "size"),
                    **{m: (m, "mean") for m in SCORES}, events=("reference_event_id", "nunique")).reset_index()
                values["site"], values["case_id"], values["population"] = site, case, population
                values["method"], values["reference_case"] = "raw", "joint_gradual"
                values["window"] = "scheduled interruption, before final channel return"
                results.append(values)
    return pd.concat(results, ignore_index=True)


def paired_days(daily, left, right, phases=None):
    part = daily if phases is None else daily.loc[daily.phase.isin(phases)]
    keys = ["day", "case_id", "horizon"]
    measures = ["n", *["sum_"+m for m in SCORES]]
    a = part.loc[part.method.eq(left)].groupby(keys)[measures].sum()
    b = part.loc[part.method.eq(right)].groupby(keys)[measures].sum()
    if not a.index.equals(b.index) or not a.n.equals(b.n):
        raise ValueError("Compared methods do not share the same daily scoring counts")
    out = a[["n"]].copy()
    for m in SCORES:
        out["left_"+m] = a["sum_"+m]
        out["right_"+m] = b["sum_"+m]
        out["delta_"+m] = a["sum_"+m]-b["sum_"+m]
    return out.reset_index()


def calendar_bootstrap(paired, start, end, block_days, replicates, seed, metric="nwis"):
    """Resample common calendar blocks and recompute equal-stratum ratios.

    Strata are case/horizon combinations with any eligible paired observations.
    A replicate missing an entire stratum is rejected and its rate reported.
    Empty calendar blocks remain in the sampling frame, preserving time gaps.
    """
    if paired.empty:
        return {"supported": False, "reason": "No eligible paired observations"}
    paired = paired.copy()
    anchor = pd.Timestamp(start).floor("D")
    last = pd.Timestamp(end).ceil("D")
    nblocks = int(np.ceil((last-anchor).days/block_days))
    paired["block"] = ((pd.to_datetime(paired.day)-anchor).dt.days//block_days).astype(int)
    if not paired.block.between(0, nblocks-1).all():
        raise ValueError("Paired dates outside the frozen calendar frame")
    strata = ["case_id", "horizon"]
    totals = paired.groupby(strata)[["n", "left_"+metric, "right_"+metric]].sum()
    index = totals.index
    matrices = {}
    for name in ("n", "left_"+metric, "right_"+metric):
        matrices[name] = paired.groupby(["block", *strata])[name].sum().unstack(strata).reindex(
            index=range(nblocks), columns=index).fillna(0.).to_numpy()
    weights = np.random.default_rng(seed).multinomial(nblocks, np.full(nblocks, 1/nblocks), size=replicates)
    with threadpool_limits(limits=1):
        counts = weights@matrices["n"]
        valid = (counts > 0).all(axis=1)
        a = (weights[valid]@matrices["left_"+metric]/counts[valid]).mean(axis=1)
        b = (weights[valid]@matrices["right_"+metric]/counts[valid]).mean(axis=1)
    point_a = float((totals["left_"+metric]/totals.n).mean())
    point_b = float((totals["right_"+metric]/totals.n).mean())
    distribution = a-b
    lo, hi = np.quantile(distribution, [.025, .975]) if len(distribution) else [np.nan]*2
    relative = np.divide(a-b, b, out=np.full(len(a), np.nan), where=b != 0)
    finite = relative[np.isfinite(relative)]
    rlo, rhi = np.quantile(finite, [.025, .975]) if len(finite) else [np.nan]*2
    return {"supported": True, "metric": metric, "left_mean": point_a, "right_mean": point_b,
        "difference": point_a-point_b, "ci_lower": lo, "ci_upper": hi,
        "relative_difference": (point_a-point_b)/point_b if point_b else np.nan,
        "relative_ci_lower": rlo, "relative_ci_upper": rhi,
        "strata": len(index), "paired_rows": int(totals.n.sum()), "eligible_days": paired.day.nunique(),
        "calendar_blocks": nblocks, "nonempty_blocks": paired.block.nunique(), "block_days": block_days,
        "bootstrap_replicates": replicates, "valid_replicates": int(valid.sum()),
        "rejected_fraction": float(1-valid.mean()), "seed": seed}


def keyed_seed(seed, *parts):
    return (seed+int(hashlib.sha256("/".join(map(str, parts)).encode()).hexdigest()[:8], 16)) % (2**32)


def collect(root, config):
    days, cells, support, diagnostics, recovery, activity = [], [], [], [], [], []
    interruption_frames = {}
    for site in config["sites"]:
        for case in config["cases"]:
            if site not in case.get("sites", config["sites"]):
                continue
            print(f"Summarizing {site}/{case['id']}", flush=True)
            scored = pd.read_parquet(root/site/f"forecasts_{case['id']}.parquet")
            scored = scored.loc[scored.valid_target & scored.daylight & scored.within_period & scored.period.eq("evaluation")].copy()
            if scored.duplicated(["origin", "horizon", "method"]).any():
                raise ValueError("Duplicate forecast keys")
            expected = scored.loc[scored.method.eq("raw"), ["origin", "horizon"]].sort_values(["origin", "horizon"]).reset_index(drop=True)
            for method, frame in scored.groupby("method"):
                keys = frame[["origin", "horizon"]].sort_values(["origin", "horizon"]).reset_index(drop=True)
                if not keys.equals(expected):
                    raise ValueError(f"Scoring-row mismatch: {site}/{case['id']}/{method}")
            scored["day"] = scored.origin.dt.floor("D")
            if case["id"] in INTERRUPTION_CASES:
                interruption_frames.setdefault(site, {})[case["id"]] = scored.loc[scored.method.eq("raw"),
                    ["origin", "horizon", "phase", "event_id", "zero_with_bright_poa", *SCORES]].copy()
            for population in ("recorded_ac", "exclude_bright_zero"):
                part = scored if population == "recorded_ac" else scored.loc[~scored.zero_with_bright_poa]
                group = ["site", "case_id", "case_group", "day", "horizon", "phase", "method"]
                daily = part.groupby(group).agg(n=("nwis", "size"), **{"sum_"+m: (m, "sum") for m in SCORES}).reset_index()
                daily["population"] = population
                days.append(daily)
                for endpoint, phases in PHASES.items():
                    selected = part if phases is None else part.loc[part.phase.isin(phases)]
                    cell = selected.groupby(["site", "case_id", "case_group", "method", "horizon"]).agg(
                        n=("nwis", "size"), **{m: (m, "mean") for m in SCORES},
                        events=("event_id", lambda x: len(set(x)-{-1})), days=("day", "nunique")).reset_index()
                    cell["population"], cell["endpoint"] = population, endpoint
                    cells.append(cell)
                    counts = cell.loc[cell.method.eq("raw")].copy()
                    counts["low_event_support"] = counts.events.lt(30) & (case["id"] != "clean")
                    counts["low_case_support"] = counts.n.lt(200)
                    support.append(counts.drop(columns=["method", *SCORES]))
            diag = pd.read_parquet(root/site/f"diagnostics_{case['id']}.parquet")
            if "label_updates_at_origin_all_horizons" in diag:
                updates = diag.loc[diag.origin.ge(pd.Timestamp(config["evaluation_start"]))].drop_duplicates(["origin", "method"])
                ua = updates.groupby("method").agg(
                    origins=("origin", "size"),
                    mean_score_updates_all_horizons=("label_updates_at_origin_all_horizons", "mean"),
                    max_score_updates_all_horizons=("label_updates_at_origin_all_horizons", "max"),
                    p95_score_updates_all_horizons=("label_updates_at_origin_all_horizons", lambda x: x.quantile(.95)),
                    fraction_origins_with_updates=("label_updates_at_origin_all_horizons", lambda x: x.gt(0).mean())).reset_index()
                ua["site"], ua["case_id"], ua["case_group"] = site, case["id"], case["group"]
                activity.append(ua)
            matched = scored.merge(diag, on=["origin", "horizon", "method"], validate="one_to_one")
            matched["unsupported"] = matched.fallback.eq("insufficient_scores")
            matched["no_matching_state"] = matched.fallback.eq("global_pool")
            ds = matched.groupby(["site", "case_id", "case_group", "phase", "method"]).agg(
                n=("unsupported", "size"), fallback_fraction=("unsupported", "mean"),
                global_fallback_fraction=("no_matching_state", "mean"),
                mean_pool=("pool_size", "mean"), mean_effective_n=("effective_n", "mean"),
                mean_score_age_h=("freshest_score_age_h", "mean"),
                nesting_repaired_fraction=("nesting_repaired", "mean")).reset_index()
            diagnostics.append(ds)
            if case["group"] == "primary" and case["id"] != "clean":
                rec = scored.loc[scored.recovery_hour.ge(0)].copy()
                rec["recovery_hour_bin"] = np.floor(rec.recovery_hour).astype(int)
                trajectory = rec.groupby(["site", "case_id", "method", "horizon", "recovery_hour_bin"]).agg(
                    n=("nwis", "size"), nwis=("nwis", "mean"), coverage90=("coverage90", "mean"),
                    events=("event_id", "nunique")).reset_index()
                recovery.append(trajectory)
    return {"daily": pd.concat(days, ignore_index=True), "cells": pd.concat(cells, ignore_index=True),
        "support": pd.concat(support, ignore_index=True), "diagnostics": pd.concat(diagnostics, ignore_index=True),
        "recovery": pd.concat(recovery, ignore_index=True),
        "activity": pd.concat(activity, ignore_index=True) if activity else pd.DataFrame(),
        "interruption": interruption_comparison(interruption_frames)}


def inference(daily, config):
    rows = []
    contrasts = [config["inference"]["primary_contrast"], *config["inference"]["secondary_contrasts"]]
    for (site, population), data in daily.groupby(["site", "population"]):
        pools = [("faulted_primary", data.loc[data.case_group.eq("primary") & ~data.case_id.eq("clean")])]
        pools += [(case, subset) for case, subset in data.groupby("case_id")]
        for pool, part in pools:
            endpoints = ["recovery_0_6", "recovery_6_24", "recovery", "failure_recovery"] if pool == "faulted_primary" else ["all" if pool == "clean" else "recovery"]
            for endpoint in endpoints:
                for left, right in contrasts:
                    paired = paired_days(part, left, right, PHASES[endpoint])
                    for block in config["inference"]["calendar_block_days"]:
                        seed = keyed_seed(config["inference"]["bootstrap_seed"], site, population, pool, endpoint, left, right, block)
                        estimate = calendar_bootstrap(paired, config["evaluation_start"], config["evaluation_end"],
                            block, config["inference"]["bootstrap_replicates"], seed)
                        rows.append({"site": site, "population": population, "pool": pool, "endpoint": endpoint,
                            "contrast": left+"_minus_"+right, **estimate})
    return pd.DataFrame(rows)


def plot_figures(output, tables):
    sites = list(tables["Summary"].site.unique())
    if sites != ["reserve10"]:
        raise ValueError("This figure adapter expects the declared reserve10 extension")
    plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False,
                        "pdf.fonttype": 42, "ps.fonttype": 42})
    figure_root = output/"figures"
    figure_root.mkdir()
    focus = tables["Contrasts"].loc[lambda d: d.population.eq("recorded_ac") & d.pool.eq("faulted_primary") &
        d.endpoint.isin(["recovery_0_6", "recovery_6_24"]) & d.contrast.eq("physical_recovery_minus_physical_availability") & d.block_days.eq(7)]
    fig, axes = plt.subplots(1, len(sites), figsize=(4.2, 2.7), constrained_layout=True, squeeze=False)
    axes = axes.ravel()
    for ax, site in zip(axes, sites):
        group = focus.loc[focus.site.eq(site)].set_index("endpoint").reindex(["recovery_0_6", "recovery_6_24"])
        delta = group.difference.to_numpy()
        ax.vlines([0, 1], group.ci_lower, group.ci_upper, color="#246695", linewidth=2)
        ax.scatter([0, 1], delta, color="#246695", zorder=3)
        ax.axhline(0, color="gray", linestyle="--", linewidth=.8)
        ax.set(xticks=[0, 1], xticklabels=["0–6 h", "6–24 h"], title="Colorado system 10",
            ylabel="Recovery − availability NWIS", xlabel="After scheduled final channel return")
    for ext in ("pdf", "png"):
        fig.savefig(figure_root/f"primary_contrast.{ext}", dpi=200)
    plt.close(fig)
    curve = tables["InterruptionComparison"].loc[lambda d: d.population.eq("recorded_ac")]
    fig, axes = plt.subplots(1, len(sites), figsize=(4.2, 2.7), constrained_layout=True, squeeze=False)
    axes = axes.ravel()
    for ax, site in zip(axes, sites):
        for case, group in curve.loc[curve.site.eq(site)].groupby("case_id"):
            label = "clean (matched times)" if case == "clean" else case.replace("_", " ")
            ax.plot(group.horizon, group.nwis, marker="o", markersize=3, label=label)
        ax.set(title="Colorado system 10", xticks=[1, 2, 3, 4],
               xlabel="Target-end clock-hour offset", ylabel="Interruption NWIS")
    axes[-1].legend(fontsize=6)
    for ext in ("pdf", "png"):
        fig.savefig(figure_root/f"failure_by_horizon.{ext}", dpi=200)
    plt.close(fig)
    trajectories = tables["Recovery"]
    fig, axes = plt.subplots(2, len(sites), figsize=(4.2, 4.2), constrained_layout=True, sharex=True, squeeze=False)
    for j, site in enumerate(sites):
        for method in ("raw", "recency", "physical_availability", "physical_recovery"):
            subset = trajectories.loc[trajectories.site.eq(site) & trajectories.method.eq(method)]
            mean = subset.groupby("recovery_hour_bin").agg(nwis=("nwis", "mean"), n=("n", "sum"))
            axes[0, j].plot(mean.index, mean.nwis, label=method.replace("physical_", ""))
            if method == "raw":
                axes[1, j].bar(mean.index, mean.n, color="#a0bacc")
        axes[0, j].set(title="Colorado system 10", ylabel="NWIS")
        axes[1, j].set(xlabel="Hours after scheduled final channel return", ylabel="Forecast-target pairs")
    axes[0, -1].legend(fontsize=6)
    for ext in ("pdf", "png"):
        fig.savefig(figure_root/f"recovery_trajectory.{ext}", dpi=200)
    plt.close(fig)


def analyze(root, output=None):
    root = Path(root)
    manifest = json.loads((root/"manifest.json").read_text())
    verification = json.loads((root/"verification.json").read_text())
    context = json.loads((root/"extension_context.json").read_text())
    if not verification["status"].startswith("passed") or not context["original_parent_authenticated"]:
        raise ValueError("Require completed authenticated extension verification")
    if verification["manifest_sha256"] != digest(root/"manifest.json"):
        raise ValueError("Run changed since verification")
    for name, expected in verification["files_sha256"].items():
        if digest(root/name) != expected:
            raise ValueError(f"Verified extension artifact changed: {name}")
    if manifest["status"] != "complete frozen evaluation":
        raise ValueError("Analyze only a completed frozen evaluation")
    output = root/"analysis" if output is None else Path(output)
    if output.exists():
        raise FileExistsError("Analysis outputs are immutable; use a new version for corrections")
    output.mkdir(parents=True)
    config = json.loads((root/"config.json").read_text())
    collected = collect(root, config)
    collected["daily"].to_parquet(output/"daily_scores.parquet", index=False)
    contrasts = inference(collected["daily"], config)
    cells = collected["cells"]
    primary = cells.loc[cells.case_group.eq("primary")]
    summary = primary.groupby(["site", "population", "endpoint", "method"]).agg(
        **{m: (m, "mean") for m in SCORES}, forecast_target_pairs=("n", "sum"), strata=("n", "size"),
        min_events_per_stratum=("events", "min"), min_cases_per_stratum=("n", "min")).reset_index()
    # Primary recovery/failure summaries contain only supported faulted cells;
    # all-period summary includes clean as one of the eight equally weighted cases.
    for site in config["sites"]:
        capacity = json.loads((root/site/"contract_primary.json").read_text())["capacity_kw_dc"]
        mask = summary.site.eq(site)
        summary.loc[mask, "wis_kw"] = summary.loc[mask, "nwis"]*capacity
        summary.loc[mask, "width90_kw"] = summary.loc[mask, "nwidth90"]*capacity
    tables = {"Summary": summary, "Cells": cells, "Contrasts": contrasts,
        "Support": collected["support"], "Diagnostics": collected["diagnostics"], "Recovery": collected["recovery"],
        "UpdateActivity": collected["activity"], "InterruptionComparison": collected["interruption"]}
    for name, frame in tables.items():
        frame.to_csv(output/f"{name}.csv", index=False)
    plot_figures(output, tables)
    workbook = Workbook()
    note = workbook.active
    note.title = "Readme"
    for row in [
        ["Run", str(root)], ["Protocol", "research/PRODUCING_SYSTEM_EXTENSION_PROTOCOL.md"],
        ["Evidence", "Separate system-10 replication after original-result exposure; 2016 training, 2017 recorded-AC evaluation; shared Colorado climate"],
        ["Weighting", "Equal supported case/horizon means; common scoring rows per method"],
        ["Uncertainty", "Paired common-calendar blocks; all replays of a weather day move together; unadjusted exploratory sensitivities"],
        ["Population", "recorded_ac retains zeros; exclude_bright_zero removes flagged scoring targets only"],
        ["Counts", "Forecast-target pairs overlap; they are not independent samples"],
        ["Guarantees", "No distribution-free or site-population guarantee; simulated receipts"],
        ["Figure source data", "Contrasts, InterruptionComparison and Recovery sheets; exact per-day sums in daily_scores.parquet"],
        ["Interruption curve", "Raw forecasts; isolated and joint loss plus clean on identical joint-gradual interruption origin/horizon rows"],
        ["Phase clock", "Scoring phases use scheduled channel returns; calibration recovery age uses observed currentness"]]:
        note.append(row)
    for name, frame in tables.items():
        sheet = workbook.create_sheet(name)
        for row in dataframe_to_rows(frame, index=False, header=True):
            sheet.append([None if isinstance(v, float) and not np.isfinite(v) else v for v in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output/"Final_Evaluation_Evidence.xlsx")
    saved = load_workbook(output/"Final_Evaluation_Evidence.xlsx", read_only=True)
    assert saved["Contrasts"].max_row == len(contrasts)+1
    saved.close()
    report = {"status": "complete analysis", "source_manifest_sha256": digest(root/"manifest.json"),
        "analysis_code_sha256": digest(Path(__file__)), "config_sha256": digest(root/"config.json"),
        "tables": {k: len(v) for k, v in tables.items()}, "daily_rows": len(collected["daily"]),
        "files_sha256": {str(p.relative_to(output)): digest(p) for p in output.rglob("*") if p.is_file()},
        "seed_derivation": "Frozen seed plus first 32 SHA256 bits of ordered site/population/pool/endpoint/contrast/block key, modulo 2^32",
        "bootstrap_missing_stratum_policy": "Reject replicate and report rate; point estimate uses supported fixed strata"}
    (output/"manifest.json").write_text(json.dumps(report, indent=2)+"\n")
    print(f"Completed analysis: {output}", flush=True)


def reproduce_tables(root, output):
    """Recompute from the exact verified archive into a fresh directory."""
    root, output = Path(root), Path(output)
    if output.exists():
        raise FileExistsError("Table reproduction requires a fresh output directory")
    verification = json.loads((root/"verification.json").read_text())
    original = json.loads((root/"analysis/manifest.json").read_text())
    if not verification["status"].startswith("passed") or original["status"] != "complete analysis":
        raise ValueError("Require a completed verified source analysis")
    checksum = digest(root/"manifest.json")
    if verification["manifest_sha256"] != checksum or original["source_manifest_sha256"] != checksum:
        raise ValueError("Source run differs from its verification or analysis")
    if original["analysis_code_sha256"] != digest(Path(__file__)):
        raise ValueError("Use the exact analysis source recorded in the original manifest")
    if not verification.get("files_sha256"):
        raise ValueError("Source verification must identify its checked file snapshot")
    for directory, checksums in ((root, verification["files_sha256"]), (root/"analysis", original["files_sha256"])):
        for name, expected in checksums.items():
            if digest(directory/name) != expected:
                raise ValueError(f"Changed source artifact: {directory/name}")
    analyze(root, output)
    comparisons = {}
    for name in original["tables"]:
        first, second = pd.read_csv(root/"analysis"/f"{name}.csv"), pd.read_csv(output/f"{name}.csv")
        pd.testing.assert_frame_equal(first, second, check_exact=False, rtol=1e-12, atol=1e-12)
        comparisons[name] = {"rows": len(first), "columns": len(first.columns), "matched": True,
            "original_sha256": digest(root/"analysis"/f"{name}.csv"), "reproduced_sha256": digest(output/f"{name}.csv")}
    result = {"status": "all principal tables reproduced from verified archived forecasts",
        "original_analysis_manifest_sha256": digest(root/"analysis/manifest.json"), "comparisons": comparisons,
        "tolerance": "DataFrame structure and values; floating comparisons rtol=1e-12, atol=1e-12",
        "scope": "Recomputes scores aggregated from archived issued forecasts and declared bootstrap; no model refitting and no claim of independent human reproduction"}
    (output/"reproduction_comparison.json").write_text(json.dumps(result, indent=2)+"\n")
    print(f"Reproduced and compared {len(comparisons)} principal tables", flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run")
    parser.add_argument("--output", help="Reproduce an existing verified analysis in a fresh directory and compare its tables")
    args = parser.parse_args()
    if args.output:
        reproduce_tables(args.run, args.output)
    else:
        analyze(args.run)
